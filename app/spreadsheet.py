"""Leitura e validacao de planilhas (xlsx/csv) para o app.

As colunas sao definidas em COLUMN_SPECS: cada uma tem o rotulo usado no modelo,
os cabecalhos alternativos aceitos na leitura e se e fixa (sempre no modelo).
Grupo, template e proxy NAO vem da planilha - sao escolhidos na tela do app.
"""

import csv
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

# Permite importar modulos da raiz do projeto (xlsx_to_csv, zabbix_*).
ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402

import zabbix_web_batch_import as core  # noqa: E402
from xlsx_to_csv import normalize_for_match, normalize_vendor  # noqa: E402

IP_RE = re.compile(r"^\d{1,3}(\.\d{1,3}){3}$")
MAC_RE = re.compile(r"^([0-9A-Fa-f]{2}[:-]){5}[0-9A-Fa-f]{2}$")


@dataclass(frozen=True)
class ColumnSpec:
    """Definicao de uma coluna da planilha."""

    key: str
    label: str
    aliases: tuple[str, ...] = ()
    sample: str = ""
    hint: str = ""
    required_value: bool = False
    fixed: bool = False


COLUMN_SPECS: tuple[ColumnSpec, ...] = (
    ColumnSpec(
        key="Name",
        label="Nome do host",
        aliases=("Name", "Nome"),
        sample="CAM-TESTE-01",
        hint="Nome da camera (obrigatorio).",
        required_value=True,
        fixed=True,
    ),
    ColumnSpec(
        key="IP",
        label="IP",
        aliases=("IP/Nome", "IP da camera"),
        sample="10.0.0.10",
        hint="IP da camera (obrigatorio).",
        required_value=True,
        fixed=True,
    ),
    ColumnSpec(
        key="Vendor",
        label="Fabricante",
        aliases=("Vendor", "Fabricante:"),
        sample="Hikvision",
        hint="Marca da camera (inventario).",
    ),
    ColumnSpec(
        key="Model",
        label="Modelo",
        aliases=("Model",),
        sample="DS-2CD2043G2-I",
        hint="Modelo da camera (inventario).",
    ),
    ColumnSpec(
        key="Firmware",
        label="Firmware",
        sample="2.840.0000000.28.R",
        hint="Versao de firmware (inventario).",
    ),
    ColumnSpec(
        key="MAC address",
        label="Endereço MAC",
        aliases=("MAC address", "MAC"),
        sample="AA-BB-CC-DD-EE-FF",
        hint="Endereco MAC (inventario).",
    ),
    ColumnSpec(
        key="Unidade",
        label="Unidade",
        aliases=("Setor", "Local"),
        sample="AME-STP2",
        hint="Unidade/local. Apenas orientacao, nao vai para o Zabbix.",
    ),
    ColumnSpec(
        key="Tag",
        label="Etiqueta",
        aliases=("Tag", "Tags"),
        sample="site:MATRIZ",
        hint="Etiqueta no formato chave:valor (varias separadas por ';').",
    ),
    ColumnSpec(
        key="Description",
        label="Descrição",
        aliases=("Descricao", "Description"),
        sample="Camera do portao principal",
        hint="Descricao do host no Zabbix.",
    ),
)

COLUMNS_BY_KEY = {spec.key: spec for spec in COLUMN_SPECS}
TARGET_FIELDS = [spec.key for spec in COLUMN_SPECS]
FIXED_COLUMNS = [spec.key for spec in COLUMN_SPECS if spec.fixed]
OPTIONAL_COLUMNS = [spec.key for spec in COLUMN_SPECS if not spec.fixed]
DEFAULT_TEMPLATE_COLUMNS = list(TARGET_FIELDS)


@dataclass
class RowResult:
    """Uma linha da planilha apos leitura, normalizacao e validacao."""

    row_number: int
    name_original: str
    name_normalized: str
    cells: dict[str, str]
    issues: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def valid(self) -> bool:
        return not self.issues


@dataclass
class SheetResult:
    filename: str
    rows: list[RowResult]
    columns_found: list[str] = field(default_factory=list)
    file_warnings: list[str] = field(default_factory=list)
    skipped_empty: int = 0
    skipped_trocas: int = 0

    @property
    def valid_rows(self) -> list[RowResult]:
        return [row for row in self.rows if row.valid]

    @property
    def rows_with_errors(self) -> list[RowResult]:
        return [row for row in self.rows if not row.valid]


def template_columns(selected_optional: list[str] | None = None) -> list[str]:
    """Colunas do modelo: fixas + opcionais escolhidas (ordem canonica)."""
    optional = OPTIONAL_COLUMNS if selected_optional is None else selected_optional
    chosen = set(FIXED_COLUMNS) | set(optional)
    return [key for key in TARGET_FIELDS if key in chosen]


def writable_optional_columns() -> list[ColumnSpec]:
    """Colunas opcionais oferecidas no popup de geracao do modelo."""
    return [COLUMNS_BY_KEY[key] for key in OPTIONAL_COLUMNS]


def parse_tags(text: str) -> list[tuple[str, str]]:
    """Converte 'chave:valor; chave2:valor2' em lista de tuplas."""
    tags: list[tuple[str, str]] = []
    for part in (text or "").split(";"):
        piece = part.strip()
        if not piece:
            continue
        if ":" in piece:
            key, value = piece.split(":", 1)
            tags.append((key.strip(), value.strip()))
        else:
            tags.append((piece, ""))
    return tags


def _normalize_header(value: str) -> str:
    """Cabecalho sem acentos e sem espacos duplos, caixa baixa."""
    return re.sub(r"\s+", " ", normalize_for_match(value)).strip()


def _alias_candidates(spec: ColumnSpec) -> list[str]:
    """Nomes de coluna aceitos para um campo destino (rotulo, chave e alias)."""
    return [spec.label, spec.key, *spec.aliases]


def _map_headers(header_cells: list[str]) -> tuple[dict[str, int], list[str]]:
    """Mapeia cabecalhos para campos destino. Retorna (indices, colunas ausentes)."""
    wanted: dict[str, list[str]] = {
        spec.key: [_normalize_header(alias) for alias in _alias_candidates(spec)]
        for spec in COLUMN_SPECS
    }
    normalized_headers = [_normalize_header(cell) for cell in header_cells]

    index_by_target: dict[str, int] = {}
    for key, aliases in wanted.items():
        for header_index, header in enumerate(normalized_headers):
            if header in aliases:
                index_by_target[key] = header_index
                break

    missing = [key for key in TARGET_FIELDS if key not in index_by_target]
    return index_by_target, missing


def _row_to_cells(
    values: list[str],
    index_by_target: dict[str, int],
) -> dict[str, str]:
    cells: dict[str, str] = {}
    for key, idx in index_by_target.items():
        cells[key] = values[idx].strip() if idx < len(values) else ""
    for key in TARGET_FIELDS:
        cells.setdefault(key, "")
    return cells


def _validate_row(row_number: int, cells: dict[str, str], dup_lines: dict[str, list[int]]) -> RowResult:
    name_original = cells.get("Name", "")
    name_normalized = core.normalize_zabbix_name(name_original)
    result = RowResult(
        row_number=row_number,
        name_original=name_original,
        name_normalized=name_normalized,
        cells=cells,
    )

    if name_original and name_normalized != name_original:
        result.warnings.append(
            f"Nome normalizado: '{name_original}' -> '{name_normalized}'"
        )

    if not name_normalized:
        result.issues.append("Nome sem caracteres validos para o Zabbix.")
    elif len(dup_lines.get(name_normalized, [])) > 1:
        other_lines = [line for line in dup_lines[name_normalized] if line != row_number]
        result.issues.append(
            "Nome duplicado apos normalizacao com a(s) linha(s): "
            + ", ".join(str(line) for line in other_lines)
        )

    ip = cells.get("IP", "")
    if not ip:
        result.issues.append("IP vazio.")
    elif not IP_RE.match(ip):
        result.issues.append(f"IP com formato invalido: '{ip}'.")

    mac = cells.get("MAC address", "")
    if mac and not MAC_RE.match(mac):
        result.warnings.append(f"MAC com formato incomum: '{mac}'.")

    if cells.get("Vendor"):
        cells["Vendor"] = normalize_vendor(cells["Vendor"])

    return result


def _rows_from_table(
    filename: str,
    header_cells: list[str],
    table_rows: list[list[str]],
    file_warnings: list[str],
    skipped_empty: int,
    skipped_trocas: int,
) -> SheetResult:
    index_by_target, missing = _map_headers(header_cells)

    required_missing = [
        COLUMNS_BY_KEY[key].label
        for key in TARGET_FIELDS
        if key in missing and COLUMNS_BY_KEY[key].required_value
    ]
    if required_missing:
        raise ValueError(
            "Coluna(s) obrigatoria(s) nao encontrada(s): "
            + ", ".join(required_missing)
            + ". Confira o cabecalho ou baixe o modelo de exemplo."
        )

    # Colunas opcionais ausentes sao normais (o modelo permite escolher quais usar).

    columns_found = [key for key in TARGET_FIELDS if key in index_by_target]
    rows: list[RowResult] = []
    seen_names: dict[str, list[int]] = {}

    for row_number, values in enumerate(table_rows, start=2):
        cells = _row_to_cells(values, index_by_target)
        name_original = cells.get("Name", "")

        if not name_original:
            skipped_empty += 1
            continue
        if "troca realizada" in normalize_for_match(name_original):
            skipped_trocas += 1
            continue

        name_normalized = core.normalize_zabbix_name(name_original)
        seen_names.setdefault(name_normalized, []).append(row_number)

    for row_number, values in enumerate(table_rows, start=2):
        cells = _row_to_cells(values, index_by_target)
        name_original = cells.get("Name", "")

        if not name_original:
            continue
        if "troca realizada" in normalize_for_match(name_original):
            continue

        rows.append(_validate_row(row_number, cells, seen_names))

    return SheetResult(
        filename=filename,
        rows=rows,
        columns_found=columns_found,
        file_warnings=file_warnings,
        skipped_empty=skipped_empty,
        skipped_trocas=skipped_trocas,
    )


def read_xlsx(path: Path) -> SheetResult:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    raw_rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    if not raw_rows:
        raise ValueError("A planilha esta vazia.")

    header_cells = [str(cell).strip() if cell is not None else "" for cell in raw_rows[0]]
    table_rows = [
        [str(cell) if cell is not None else "" for cell in row]
        for row in raw_rows[1:]
    ]
    return _rows_from_table(path.name, header_cells, table_rows, [], 0, 0)


def read_csv(path: Path) -> SheetResult:
    encodings = ["utf-8-sig", "cp1252", "latin-1"]
    content: list[list[str]] = []
    used_encoding = None
    for encoding in encodings:
        try:
            with path.open("r", encoding=encoding, newline="") as file:
                content = list(csv.reader(file))
            used_encoding = encoding
            break
        except UnicodeDecodeError:
            continue
    if not content:
        raise ValueError("Nao foi possivel ler o CSV (encoding desconhecido).")

    header_cells = [cell.strip() for cell in content[0]]
    table_rows = [[cell for cell in row] for row in content[1:]]

    file_warnings: list[str] = []
    if used_encoding and used_encoding != "utf-8-sig":
        file_warnings.append(f"CSV lido como {used_encoding}.")

    return _rows_from_table(path.name, header_cells, table_rows, file_warnings, 0, 0)


def read_file(path: Path) -> SheetResult:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv(path)
    if suffix == ".xlsx":
        return read_xlsx(path)
    raise ValueError(
        "Formato nao suportado. Use .xlsx ou .csv (se tiver um .xls antigo, salve como .xlsx)."
    )


def write_example_template(path: Path, columns: list[str] | None = None) -> None:
    """Gera um xlsx de exemplo com as colunas escolhidas e uma linha de amostra."""
    selected = template_columns(columns)
    specs = [COLUMNS_BY_KEY[key] for key in selected]

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Cameras"
    sheet.append([spec.label for spec in specs])
    sheet.append([spec.sample for spec in specs])

    for column_cells in sheet.columns:
        width = max(len(str(cell.value)) for cell in column_cells if cell.value is not None) + 2
        column_letter = column_cells[0].column_letter
        sheet.column_dimensions[column_letter].width = min(width, 60)
    workbook.save(path)
    workbook.close()

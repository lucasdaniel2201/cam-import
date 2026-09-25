"""Testes de leitura/validacao de planilhas (app/spreadsheet.py).

Cobrem o mapeamento de colunas por apelido, a geracao do modelo, a validacao
linha a linha e os invariantes da definicao de colunas.
"""

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402

from app import spreadsheet as sp  # noqa: E402


def write_sheet(path: Path, header: list[str], rows: list[list[str]]) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


class TestParseTags(unittest.TestCase):
    def test_chave_valor(self):
        self.assertEqual(sp.parse_tags("site:MATRIZ"), [("site", "MATRIZ")])

    def test_multiplas_separadas_por_ponto_e_virgula(self):
        self.assertEqual(
            sp.parse_tags("site:MATRIZ; andar:2"),
            [("site", "MATRIZ"), ("andar", "2")],
        )

    def test_sem_valor(self):
        self.assertEqual(sp.parse_tags("IMPORTADO"), [("IMPORTADO", "")])

    def test_valor_com_dois_pontos(self):
        self.assertEqual(sp.parse_tags("url:http://x/y"), [("url", "http://x/y")])

    def test_espacos_sao_removidos(self):
        self.assertEqual(sp.parse_tags("  site : MATRIZ  "), [("site", "MATRIZ")])

    def test_vazio(self):
        self.assertEqual(sp.parse_tags(""), [])
        self.assertEqual(sp.parse_tags("  ;  "), [])


class TestTemplateColumns(unittest.TestCase):
    def test_padrao_inclui_todas(self):
        self.assertEqual(sp.template_columns(), sp.TARGET_FIELDS)

    def test_subconjunto_mantem_fixas_e_ordem(self):
        columns = sp.template_columns(["Tag"])
        self.assertIn("Name", columns)
        self.assertIn("IP", columns)
        self.assertIn("Tag", columns)
        self.assertNotIn("Vendor", columns)
        # ordem canonica preservada
        self.assertEqual(columns, [c for c in sp.TARGET_FIELDS if c in columns])

    def test_chave_desconhecida_e_ignorada(self):
        self.assertEqual(sp.template_columns(["Inexistente"]), ["Name", "IP"])

    def test_nao_aceita_remover_fixa(self):
        columns = sp.template_columns([])
        self.assertEqual(columns, sp.FIXED_COLUMNS)


class TestColumnSpecsIntegridade(unittest.TestCase):
    def test_chaves_unicas(self):
        keys = [spec.key for spec in sp.COLUMN_SPECS]
        self.assertEqual(len(keys), len(set(keys)))

    def test_rotulos_unicos(self):
        labels = [spec.label for spec in sp.COLUMN_SPECS]
        self.assertEqual(len(labels), len(set(labels)))

    def test_fixas_sao_apenas_name_e_ip(self):
        self.assertEqual(sp.FIXED_COLUMNS, ["Name", "IP"])

    def test_colunas_fixas_sao_obrigatorias(self):
        for key in sp.FIXED_COLUMNS:
            self.assertTrue(
                sp.COLUMNS_BY_KEY[key].required_value,
                f"{key} e fixa e deveria exigir valor",
            )

    def test_apelidos_nao_colidem_entre_colunas(self):
        """Um cabecalho nao pode apontar para duas colunas diferentes."""
        seen: dict[str, str] = {}
        for spec in sp.COLUMN_SPECS:
            for alias in (spec.label, spec.key, *spec.aliases):
                normalized = sp._normalize_header(alias)
                other = seen.get(normalized)
                # repetir o mesmo apelido dentro da propria coluna e permitido
                if other is not None and other != spec.key:
                    self.fail(
                        f"apelido '{alias}' colide entre '{other}' e '{spec.key}'"
                    )
                seen[normalized] = spec.key

    def test_sem_coluna_server(self):
        self.assertNotIn("Server", sp.TARGET_FIELDS)

    def test_grupo_template_proxy_nao_sao_colunas(self):
        for key in ("Group", "Template", "Proxy"):
            self.assertNotIn(key, sp.TARGET_FIELDS)


class TestWriteExampleTemplate(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_cabecalho_usa_rotulos_do_modelo(self):
        path = self.tmp / "modelo.xlsx"
        sp.write_example_template(path)
        header = [c.value for c in openpyxl.load_workbook(path).active[1]]
        self.assertEqual(header, [sp.COLUMNS_BY_KEY[k].label for k in sp.TARGET_FIELDS])

    def test_modelo_nao_tem_coluna_server(self):
        path = self.tmp / "modelo.xlsx"
        sp.write_example_template(path)
        header = [c.value for c in openpyxl.load_workbook(path).active[1]]
        self.assertNotIn("Server", header)

    def test_modelo_tem_etiqueta_e_descricao(self):
        path = self.tmp / "modelo.xlsx"
        sp.write_example_template(path)
        header = [c.value for c in openpyxl.load_workbook(path).active[1]]
        self.assertIn("Etiqueta", header)
        self.assertIn("Descrição", header)

    def test_subconjunto_de_colunas(self):
        path = self.tmp / "modelo.xlsx"
        sp.write_example_template(path, ["Tag"])
        header = [c.value for c in openpyxl.load_workbook(path).active[1]]
        self.assertEqual(
            header,
            [sp.COLUMNS_BY_KEY[k].label for k in ("Name", "IP", "Tag")],
        )

    def test_modelo_gerado_pode_ser_lido_e_evalido(self):
        path = self.tmp / "modelo.xlsx"
        sp.write_example_template(path)
        sheet = sp.read_file(path)
        self.assertEqual(len(sheet.rows), 1)
        self.assertEqual(len(sheet.valid_rows), 1)
        self.assertEqual(sheet.rows[0].cells["Tag"], "site:MATRIZ")


class TestReadAndValidate(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def _read(self, header: list[str], rows: list[list[str]]):
        path = self.tmp / "dados.xlsx"
        write_sheet(path, header, rows)
        return sp.read_file(path)

    def test_apelidos_aceitos(self):
        sheet = self._read(
            ["Nome", "IP", "Fabricante:", "MAC"],
            [["CAM-01", "10.0.0.1", "Hikvision", "AA-BB-CC-DD-EE-FF"]],
        )
        cells = sheet.rows[0].cells
        self.assertEqual(cells["Name"], "CAM-01")
        self.assertEqual(cells["Vendor"], "Hikvision")
        self.assertEqual(cells["MAC address"], "AA-BB-CC-DD-EE-FF")

    def test_coluna_obrigatoria_ausente(self):
        with self.assertRaises(ValueError) as ctx:
            self._read(["Fabricante"], [["Hikvision"]])
        self.assertIn("obrigatoria", str(ctx.exception))

    def test_coluna_opcional_ausente_nao_gera_aviso(self):
        sheet = self._read(["Nome do host", "IP"], [["CAM-01", "10.0.0.1"]])
        self.assertEqual(sheet.file_warnings, [])

    def test_ip_vazio_e_erro(self):
        sheet = self._read(["Nome do host", "IP"], [["CAM-01", ""]])
        self.assertFalse(sheet.rows[0].valid)
        self.assertTrue(any("IP vazio" in i for i in sheet.rows[0].issues))

    def test_ip_invalido_e_erro(self):
        sheet = self._read(["Nome do host", "IP"], [["CAM-01", "999.1.1"]])
        self.assertTrue(any("formato invalido" in i for i in sheet.rows[0].issues))

    def test_nome_duplicado_apos_normalizacao(self):
        sheet = self._read(
            ["Nome do host", "IP"],
            [["CAM-Ação", "10.0.0.1"], ["CAM-Acao", "10.0.0.2"]],
        )
        self.assertEqual(len(sheet.rows_with_errors), 2)
        self.assertTrue(all("duplicado" in i for r in sheet.rows for i in r.issues))

    def test_aviso_de_normalizacao(self):
        sheet = self._read(["Nome do host", "IP"], [["Câmera São", "10.0.0.1"]])
        row = sheet.rows[0]
        self.assertTrue(row.valid)
        self.assertEqual(row.name_normalized, "Camera Sao")
        self.assertTrue(any("normalizado" in w for w in row.warnings))

    def test_mac_estranho_gera_aviso(self):
        sheet = self._read(
            ["Nome do host", "IP", "Endereço MAC"],
            [["CAM-01", "10.0.0.1", "ZZZZ"]],
        )
        row = sheet.rows[0]
        self.assertTrue(row.valid)
        self.assertTrue(any("MAC" in w for w in row.warnings))

    def test_mac_valido_nao_gera_aviso(self):
        sheet = self._read(
            ["Nome do host", "IP", "Endereço MAC"],
            [["CAM-01", "10.0.0.1", "AA:BB:CC:DD:EE:FF"]],
        )
        self.assertEqual(sheet.rows[0].warnings, [])

    def test_linha_sem_nome_e_ignorada(self):
        sheet = self._read(
            ["Nome do host", "IP"],
            [["CAM-01", "10.0.0.1"], ["", "10.0.0.2"]],
        )
        self.assertEqual(len(sheet.rows), 1)
        self.assertEqual(sheet.skipped_empty, 1)

    def test_troca_realizada_e_ignorada(self):
        sheet = self._read(
            ["Nome do host", "IP"],
            [["CAM-01", "10.0.0.1"], ["CAM-02 - TROCA REALIZADA", "10.0.0.2"]],
        )
        self.assertEqual(len(sheet.rows), 1)
        self.assertEqual(sheet.skipped_trocas, 1)

    def test_vendor_hikvision_padronizado(self):
        sheet = self._read(
            ["Nome do host", "IP", "Fabricante"],
            [["CAM-01", "10.0.0.1", "HIKVISION"]],
        )
        self.assertEqual(sheet.rows[0].cells["Vendor"], "Hikvision")

    def test_csv_tambem_e_lido(self):
        path = self.tmp / "dados.csv"
        path.write_text("Nome do host,IP\nCAM-01,10.0.0.1\n", encoding="utf-8")
        sheet = sp.read_file(path)
        self.assertEqual(len(sheet.valid_rows), 1)

    def test_formato_nao_suportado(self):
        path = self.tmp / "dados.txt"
        path.write_text("x", encoding="utf-8")
        with self.assertRaises(ValueError):
            sp.read_file(path)


if __name__ == "__main__":
    unittest.main(verbosity=2)
